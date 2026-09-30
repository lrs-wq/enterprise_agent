import csv
import json
import platform
import re
import shutil
import subprocess
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import TypedDict
from zipfile import BadZipFile, ZipFile

import xlrd
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".xls", ".csv", ".txt", ".md"}
MAX_CHUNK_CHARS = 1_400
CHUNK_OVERLAP_CHARS = 160
PROJECT_ROOT = Path(__file__).resolve().parents[2]
MACOS_OCR_SCRIPT = PROJECT_ROOT / "scripts" / "pdf_ocr_vision.swift"
WPS_DISPIMG_PATTERN = re.compile(r'^=?DISPIMG\(\s*"[^"]+"\s*,\s*\d+\s*\)$', re.IGNORECASE)


class OcrPage(TypedDict):
    page: int
    text: str


@dataclass(frozen=True)
class ParsedChunk:
    content: str
    page: int | None = None
    sheet: str | None = None
    row_start: int | None = None
    row_end: int | None = None


def normalize_text(value: object) -> str:
    text = str(value).replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def split_text(text: str, max_chars: int = MAX_CHUNK_CHARS) -> list[str]:
    text = normalize_text(text)
    if not text:
        return []
    if len(text) <= max_chars:
        return [text]

    paragraphs = [item.strip() for item in re.split(r"\n\s*\n", text) if item.strip()]
    if len(paragraphs) == 1:
        paragraphs = [
            item.strip() for item in re.split(r"(?<=[。！？；.!?;])\s*", text) if item.strip()
        ]

    chunks: list[str] = []
    current = ""
    for paragraph in paragraphs:
        if len(paragraph) > max_chars:
            if current:
                chunks.append(current.strip())
                current = ""
            start = 0
            while start < len(paragraph):
                end = min(start + max_chars, len(paragraph))
                chunks.append(paragraph[start:end].strip())
                if end == len(paragraph):
                    break
                start = max(end - CHUNK_OVERLAP_CHARS, start + 1)
            continue

        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= max_chars:
            current = candidate
            continue

        chunks.append(current.strip())
        overlap = current[-CHUNK_OVERLAP_CHARS:].strip()
        current = f"{overlap}\n\n{paragraph}".strip() if overlap else paragraph

    if current:
        chunks.append(current.strip())
    return [chunk for chunk in chunks if chunk]


def parse_pdf(
    path: Path,
    *,
    max_pages: int = 300,
    ocr_timeout_seconds: int = 240,
) -> list[ParsedChunk]:
    reader = PdfReader(str(path))
    if len(reader.pages) > max_pages:
        raise ValueError(f"PDF 页数超过 {max_pages} 页限制")
    chunks: list[ParsedChunk] = []
    missing_pages: set[int] = set()
    for page_number, page in enumerate(reader.pages, start=1):
        text = normalize_text(page.extract_text() or "")
        if not text:
            missing_pages.add(page_number)
            continue
        for content in split_text(text):
            chunks.append(ParsedChunk(content=content, page=page_number))
    if missing_pages:
        try:
            chunks.extend(
                parse_pdf_with_macos_vision(
                    path,
                    page_numbers=missing_pages,
                    timeout_seconds=ocr_timeout_seconds,
                )
            )
        except ValueError:
            if not chunks:
                raise
    chunks.sort(key=lambda chunk: chunk.page or 0)
    return chunks


def parse_pdf_with_macos_vision(
    path: Path,
    page_numbers: set[int] | None = None,
    *,
    timeout_seconds: int = 240,
) -> list[ParsedChunk]:
    """Use Apple's on-device Vision OCR as a local fallback for scanned PDFs."""
    xcrun = shutil.which("xcrun")
    if platform.system() != "Darwin" or xcrun is None or not MACOS_OCR_SCRIPT.exists():
        raise ValueError("PDF 中没有可提取文本；当前系统未配置扫描版 PDF OCR，请先进行 OCR")

    try:
        completed = subprocess.run(
            [xcrun, "swift", str(MACOS_OCR_SCRIPT), str(path)],
            check=True,
            capture_output=True,
            text=True,
            timeout=timeout_seconds,
        )
        pages: list[OcrPage] = json.loads(completed.stdout)
    except subprocess.TimeoutExpired as exc:
        raise ValueError("扫描版 PDF OCR 超时") from exc
    except (subprocess.CalledProcessError, json.JSONDecodeError) as exc:
        stderr = getattr(exc, "stderr", "") or ""
        detail = normalize_text(stderr)[-500:]
        message = f"扫描版 PDF OCR 失败：{detail}" if detail else "扫描版 PDF OCR 失败"
        raise ValueError(message) from exc

    chunks: list[ParsedChunk] = []
    for item in pages:
        text = normalize_text(item.get("text", ""))
        page_number = item.get("page")
        if not text or not isinstance(page_number, int):
            continue
        if page_numbers is not None and page_number not in page_numbers:
            continue
        for content in split_text(text):
            chunks.append(ParsedChunk(content=content, page=page_number))
    if not chunks:
        raise ValueError("扫描版 PDF 已完成 OCR，但没有识别到可用文字")
    return chunks


def row_to_text(headers: list[str], values: Iterable[object]) -> str:
    pairs: list[str] = []
    for index, raw_value in enumerate(values):
        if raw_value is None:
            continue
        value = normalize_text(raw_value)
        if not value:
            continue
        if WPS_DISPIMG_PATTERN.fullmatch(value):
            continue
        header = headers[index] if index < len(headers) and headers[index] else f"第{index + 1}列"
        pairs.append(f"{header}: {value}")
    return "\n".join(pairs)


def make_headers(values: Iterable[object]) -> list[str]:
    headers: list[str] = []
    for index, value in enumerate(values):
        normalized = normalize_text(value) if value is not None else ""
        headers.append(normalized or f"第{index + 1}列")
    return headers


def validate_office_archive(
    path: Path,
    *,
    max_uncompressed_bytes: int,
    max_members: int,
) -> None:
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
    except BadZipFile as exc:
        raise ValueError("Excel 文件压缩结构损坏") from exc
    if len(members) > max_members:
        raise ValueError(f"Excel 压缩成员超过 {max_members} 个限制")
    unpacked = sum(member.file_size for member in members)
    if unpacked > max_uncompressed_bytes:
        limit_mb = max_uncompressed_bytes // (1024 * 1024)
        raise ValueError(f"Excel 解压后体积超过 {limit_mb}MB 限制")


def parse_xlsx(
    path: Path,
    *,
    max_rows: int = 50_000,
    max_archive_uncompressed_bytes: int = 256 * 1024 * 1024,
    max_archive_members: int = 5000,
) -> list[ParsedChunk]:
    validate_office_archive(
        path,
        max_uncompressed_bytes=max_archive_uncompressed_bytes,
        max_members=max_archive_members,
    )
    workbook = load_workbook(path, read_only=True, data_only=True)
    chunks: list[ParsedChunk] = []
    rows_seen = 0
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            first_row = next(rows, None)
            if first_row is None:
                continue
            rows_seen += 1
            headers = make_headers(first_row)
            for row_number, values in enumerate(rows, start=2):
                rows_seen += 1
                if rows_seen > max_rows:
                    raise ValueError(f"Excel 数据行超过 {max_rows} 行限制")
                content = row_to_text(headers, values)
                if not content:
                    continue
                prefixed = f"工作表: {worksheet.title}\n行号: {row_number}\n{content}"
                for part in split_text(prefixed):
                    chunks.append(
                        ParsedChunk(
                            content=part,
                            sheet=worksheet.title,
                            row_start=row_number,
                            row_end=row_number,
                        )
                    )
    finally:
        workbook.close()
    if not chunks:
        raise ValueError("Excel 中没有可导入的数据行")
    return chunks


def parse_xls(path: Path, *, max_rows: int = 50_000) -> list[ParsedChunk]:
    workbook = xlrd.open_workbook(str(path), on_demand=True)
    chunks: list[ParsedChunk] = []
    rows_seen = 0
    try:
        for sheet in workbook.sheets():
            if sheet.nrows == 0:
                continue
            rows_seen += sheet.nrows
            if rows_seen > max_rows:
                raise ValueError(f"Excel 数据行超过 {max_rows} 行限制")
            headers = make_headers(sheet.row_values(0))
            for row_index in range(1, sheet.nrows):
                row_number = row_index + 1
                content = row_to_text(headers, sheet.row_values(row_index))
                if not content:
                    continue
                prefixed = f"工作表: {sheet.name}\n行号: {row_number}\n{content}"
                for part in split_text(prefixed):
                    chunks.append(
                        ParsedChunk(
                            content=part,
                            sheet=sheet.name,
                            row_start=row_number,
                            row_end=row_number,
                        )
                    )
    finally:
        workbook.release_resources()
    if not chunks:
        raise ValueError("Excel 中没有可导入的数据行")
    return chunks


def parse_csv(path: Path, *, max_rows: int = 50_000) -> list[ParsedChunk]:
    chunks: list[ParsedChunk] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        first_row = next(reader, None)
        if first_row is None:
            raise ValueError("CSV 文件为空")
        headers = make_headers(first_row)
        for row_number, values in enumerate(reader, start=2):
            if row_number > max_rows:
                raise ValueError(f"CSV 数据行超过 {max_rows} 行限制")
            content = row_to_text(headers, values)
            if content:
                chunks.append(
                    ParsedChunk(content=content, row_start=row_number, row_end=row_number)
                )
    if not chunks:
        raise ValueError("CSV 中没有可导入的数据行")
    return chunks


def parse_text(path: Path) -> list[ParsedChunk]:
    text = path.read_text(encoding="utf-8-sig")
    chunks = [ParsedChunk(content=content) for content in split_text(text)]
    if not chunks:
        raise ValueError("文本文件为空")
    return chunks


def parse_document(
    path: Path,
    *,
    max_pdf_pages: int = 300,
    ocr_timeout_seconds: int = 240,
    max_spreadsheet_rows: int = 50_000,
    max_archive_uncompressed_bytes: int = 256 * 1024 * 1024,
    max_archive_members: int = 5000,
) -> list[ParsedChunk]:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的文件类型: {extension}")
    if extension == ".pdf":
        return parse_pdf(
            path,
            max_pages=max_pdf_pages,
            ocr_timeout_seconds=ocr_timeout_seconds,
        )
    if extension == ".xlsx":
        return parse_xlsx(
            path,
            max_rows=max_spreadsheet_rows,
            max_archive_uncompressed_bytes=max_archive_uncompressed_bytes,
            max_archive_members=max_archive_members,
        )
    if extension == ".xls":
        return parse_xls(path, max_rows=max_spreadsheet_rows)
    if extension == ".csv":
        return parse_csv(path, max_rows=max_spreadsheet_rows)
    return parse_text(path)
